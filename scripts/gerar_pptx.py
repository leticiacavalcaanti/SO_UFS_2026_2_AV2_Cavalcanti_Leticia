#!/usr/bin/env python3
"""Gera report/apresentacao.pptx (editável) com os mesmos 15 slides de apresentacao.md.

Requer python-pptx e Pillow:  pip install python-pptx pillow
Uso (a partir da raiz do repositório):  python scripts/gerar_pptx.py
"""
import os
import re

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
RES = os.path.join(ROOT, "results")
OUT = os.path.join(ROOT, "report", "apresentacao.pptx")

NAVY, BLUE, ORANGE = "0B2545", "2A78D6", "EB6834"
INK, MUTED, TINT, WHITE = "1A1A1A", "52514E", "EEF3FA", "FFFFFF"
FONT = "Calibri"
TOTAL = 15


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
LAYOUT_TITLE_ONLY = prs.slide_layouts[5]
state = {"n": 0}


def style_run(r, size, bold=False, color=INK, italic=False, font=FONT):
    r.font.size, r.font.bold, r.font.italic, r.font.name = Pt(size), bold, italic, font
    r.font.color.rgb = rgb(color)


def add_runs(p, text, size, color=INK, bold=False):
    """Texto com **negrito**, *itálico* e `código` inline."""
    for part in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*|`.+?`)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run()
            r.text = part[2:-2]
            style_run(r, size, True, color)
        elif part.startswith("*"):
            r = p.add_run()
            r.text = part[1:-1]
            style_run(r, size, bold, color, italic=True)
        elif part.startswith("`"):
            r = p.add_run()
            r.text = part[1:-1]
            style_run(r, size, bold, color, font="Consolas")
        else:
            r = p.add_run()
            r.text = part
            style_run(r, size, bold, color)


def new_slide(title, dark=False):
    s = prs.slides.add_slide(LAYOUT_TITLE_ONLY)
    state["n"] += 1
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(NAVY if dark else WHITE)
    t = s.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0)
    tf = t.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = title
    style_run(r, 32, True, WHITE if dark else NAVY)
    if not dark:
        n = s.shapes.add_textbox(Inches(11.6), Inches(6.95), Inches(1.2), Inches(0.35))
        n.name = "Numero do slide"
        q = n.text_frame.paragraphs[0]
        q.alignment = PP_ALIGN.RIGHT
        r = q.add_run()
        r.text = f"{state['n']} / {TOTAL}"
        style_run(r, 11, False, MUTED)
    return s


def bullets(s, items, x, y, w, h, size=18, color=INK, name="Texto", bullet=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        if bullet:
            pPr.set("marL", str(Inches(0.3)))
            pPr.set("indent", str(-Inches(0.3)))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
        else:
            pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        p.space_after = Pt(6)
        add_runs(p, it, size, color)
    return tb


def picture(s, fname, x, y, w):
    path = os.path.join(RES, fname)
    with Image.open(path) as im:
        h = w * im.height / im.width
    pic = s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
    pic.name = f"Figura {fname}"
    return h


def table(s, rows, x, y, w, col_w, size=16, row_h=0.5):
    nr, nc = len(rows), len(rows[0])
    gt = s.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    gt.name = "Tabela"
    tbl = gt.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        tbl.rows[i].height = Inches(row_h)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(NAVY if i == 0 else (TINT if i % 2 else WHITE))
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            add_runs(p, val, size, WHITE if i == 0 else INK, bold=(i == 0))
    return tbl


def box(s, text, x, y, w, h, fill=TINT, color=INK, size=16, bold=False, name="Caixa", shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.name = name
    b.fill.solid()
    b.fill.fore_color.rgb = rgb(fill)
    b.line.fill.background()
    b.shadow.inherit = False
    tf = b.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        add_runs(p, line, size, color, bold)
    return b


def arrow(s, x1, y1, x2, y2):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(MUTED)
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(qn("a:tailEnd"), {"type": "triangle"}))
    return c


# ---------------------------------------------------------------- 1 capa
s = new_slide("GenResearch-Servidor", dark=True)
s.shapes.title.top, s.shapes.title.height = Inches(2.2), Inches(1.2)
s.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(48)
bullets(s, ["Memória, armazenamento, sincronização e desempenho em um serviço RAG"], 0.6, 3.5, 12, 0.8, 26, WHITE, "Subtitulo", bullet=False)
bullets(s, ["Sistemas Operacionais — AV2 — 2026.2 · Leticia Cavalcanti (individual)",
            "*Como memória, armazenamento, caches, filas e sincronização afetam desempenho, estabilidade e correção de uma aplicação de IA generativa/RAG?*"],
        0.6, 4.6, 12, 1.8, 18, "CADCFC", "Autoria e pergunta", bullet=False)

# ---------------------------------------------------------------- 2 problema
s = new_slide("Problema e recorte")
bullets(s, ["RAG mobiliza **arquivos grandes**, **memória virtual**, **cache de páginas**, **filas**, **caches** e **logs** concorrentes",
            "Foco: **o que o SO faz**, não a qualidade da resposta",
            "Recorte: serviço próprio, **vetores e geração simulados**, corpus **sintético** (sem chave de API, sem pesos, reprodutível)",
            "Configurações **C1–C6** + Parte D (sincronização) + Parte E (memória virtual), 3 repetições cada"],
        0.6, 1.7, 12.1, 4.8, 22)

# ---------------------------------------------------------------- 3 ambiente
s = new_slide("Ambiente")
table(s, [["Item", "Valor"],
          ["SO", "Debian 13 em **contêiner** (Docker Desktop, WSL2)"],
          ["Kernel", "6.6.87.2-microsoft-standard-WSL2"],
          ["CPU / RAM / swap", "i5-10210U, 4 CPUs lógicas · 5,8 GiB · 1 GiB"],
          ["GPU", "nenhuma (somente CPU)"],
          ["Pilha", "Python 3.12 · numpy 2.1.3 · SQLite"]],
      0.6, 1.6, 12.1, [3.4, 8.7], 20, 0.62)
bullets(s, ["Comandos: `uname`, `lscpu`, `free`, `/proc/meminfo`, `lsblk`, `df`, `vmstat`, `iostat` → `results/env/`"],
        0.6, 5.7, 12.1, 0.9, 16, MUTED)

# ---------------------------------------------------------------- 4 arquitetura
s = new_slide("Arquitetura")
y1, y2, bh = 1.7, 3.55, 1.1
box(s, "**2 produtores**", 0.6, y1, 2.6, bh, name="Produtores")
box(s, "**BoundedBuffer**\nempty / full / mutex", 4.1, y1, 3.3, bh, name="Fila")
box(s, "**N threads**\ntrabalhadoras", 8.3, y1, 2.6, bh, name="Trabalhadores")
arrow(s, 3.2, y1 + bh / 2, 4.1, y1 + bh / 2)
arrow(s, 7.4, y1 + bh / 2, 8.3, y1 + bh / 2)
arrow(s, 9.6, y1 + bh, 9.6, y2)
box(s, "**Cache** LRU/FIFO\n+ mutex", 8.3, y2, 2.6, bh, fill="DCE9FB", name="Cache")
box(s, "Embedding +\n**varredura do índice**\n(.npy: read | mmap)", 4.1, y2, 3.3, bh, fill="DCE9FB", name="Indice")
box(s, "**SQLite** → contexto\n→ geração simulada", 0.6, y2, 2.9, bh, fill="DCE9FB", name="Geracao")
arrow(s, 8.3, y2 + bh / 2, 7.4, y2 + bh / 2)
arrow(s, 4.1, y2 + bh / 2, 3.5, y2 + bh / 2)
bullets(s, ["Threads: carga **mista** (sleep, numpy e hashlib liberam o GIL)",
            "Recursos compartilhados: **fila, cache, log (mutex), conexões SQLite**"],
        0.6, 5.2, 12.1, 1.5, 20)

# ---------------------------------------------------------------- 5 C1xC2
s = new_slide("C1 × C2: read() × mmap (768 MB)")
h = picture(s, "fig_c1c2_leitura.png", 1.15, 1.4, 11.0)
bullets(s, ["**2ª execução 2–3× mais rápida**, 0 MB lidos do disco → **cache de páginas**",
            "`read()` do arquivo inteiro: **+768 MB de RSS**, 2× mais lento que ler em blocos",
            "`mmap`: VSZ sobe na abertura, RSS só ao tocar; 6.146 faltas maiores (~125 KiB cada, *readahead*)"],
        0.6, 1.4 + h + 0.15, 12.1, 1.9, 17)

# ---------------------------------------------------------------- 6 aleatorio
s = new_slide("mmap em acesso aleatório frio")
h = picture(s, "fig_c1c2_aleatorio.png", 0.6, 1.6, 7.2)
box(s, "**23 MB**\nlidos do disco por `pread()`", 8.4, 1.6, 4.3, 1.35, fill=TINT, size=20, name="Destaque pread")
box(s, "**287 MB**\nlidos do disco por `mmap`", 8.4, 3.15, 4.3, 1.35, fill="FBE3D9", size=20, name="Destaque mmap")
bullets(s, ["3.000 acessos de 4 KiB (≈ 12 MB úteis): `mmap` com **12× de amplificação** e 2,2× mais lento",
            "Cache quente: diferença pequena (12,8 × 17,6 ms)",
            "→ *mmap não é vantajoso para acesso esparso e frio*"],
        0.6, 5.0 if h < 3.4 else 5.2, 12.1, 1.8, 18)

# ---------------------------------------------------------------- 7 C5 limite
s = new_slide("C5: índice maior que a memória")
h = picture(s, "fig_c5_limite.png", 0.6, 1.5, 6.4)
bullets(s, ["Índice de 1 GB, limite de **384 MB**",
            "`read` → **OOM kill (exit 137)** nas 3 repetições",
            "`mmap` → **sobrevive**, mas com **thrashing**: ~125 mil faltas maiores, **15,3 GB lidos**, consulta **13–17× mais lenta**",
            "Swap **não** foi exercitado (contêiner sem swap)"],
        7.3, 1.5, 5.5, 4.8, 20)

# ---------------------------------------------------------------- 8 C5 escala
s = new_slide("C5: crescimento do corpus e do contexto")
h = picture(s, "fig_c5_escala.png", 1.15, 1.4, 11.0)
bullets(s, ["Recuperação **linear** no índice (≈ 70 µs/MB); RSS ≈ índice + ~40 MB",
            "`read` e `mmap` iguais em regime estacionário: muda **quando** o custo é pago",
            "Contexto de 500 → 8.000 B: geração de 20 → ~110 ms"],
        0.6, 1.4 + h + 0.15, 12.1, 1.9, 17)

# ---------------------------------------------------------------- 9 C3
s = new_slide("C3: concorrência e fila")
h = picture(s, "fig_c3_concorrencia.png", 1.15, 1.4, 11.0)
bullets(s, ["4 workers: **1,7–1,85×** a vazão · 8 workers: **sem ganho**, p95 **~5× pior**",
            "Fila 64: espera de **1,2–2,3 s** (fila 4: ~0,1 s) → fila grande = latência, não vazão",
            "0 perdas, 0 duplicatas, 0 violações de FIFO"],
        0.6, 1.4 + h + 0.15, 12.1, 1.9, 17)

# ---------------------------------------------------------------- 10 C4
s = new_slide("C4: cache compartilhado")
h = picture(s, "fig_c4_cache.png", 1.4, 1.4, 10.5)
bullets(s, ["90% de repetição: **89,2% de acerto**, tempo **9,9 s → 1,25 s** (≈ 8×)",
            "Consultas distintas: **0%** de acerto · LRU ≈ FIFO (repetição uniforme)"],
        0.6, 1.4 + h + 0.15, 12.1, 1.5, 18)

# ---------------------------------------------------------------- 11 sincronização
s = new_slide("Sincronização e falha controlada")
table(s, [["Demonstração", "Sem proteção", "Com proteção"],
          ["Cache: atualizações perdidas", "**86%**", "0 (mutex ou lock por chave)"],
          ["Cache: computações duplicadas", "22", "0"],
          ["Log: linhas corrompidas", "**76%**", "0 (mutex)"],
          ["Deadlock A→B / B→A (subprocesso)", "travou → **kill** após 3 s", "ordenação de locks: 0,3 s · timeout + recuo: 0,6 s"]],
      0.6, 1.6, 12.1, [4.4, 3.2, 4.5], 17, 0.7)
bullets(s, ["Threads travadas em `futex_wait_queue` · lock por chave: **−18% de tempo** vs. mutex global",
            "Janelas de corrida ampliadas de propósito (`sleep(0)`): injeção de falha controlada"],
        0.6, 5.3, 12.1, 1.5, 17)

# ---------------------------------------------------------------- 12 C6
s = new_slide("C6: falhas controladas")
table(s, [["Cenário", "Resultado"],
          ["Índice ausente", "40 erros tratados, 80 retentativas, processo estável"],
          ["Índice volta em 0,6 s", "**120/120** concluídas com retentativas"],
          ["Fila cheia", "**rejeita**: 112 de 120 descartadas · **bloqueia**: 120 ok em 4 s"],
          ["Serviço externo 30% falho", "**34% de erro → 0,8%** com 3 retentativas (+88% de tempo)"],
          ["Limite de memória", "OOM kill do `read` (exit 137)"]],
      0.6, 1.6, 12.1, [3.8, 8.3], 18, 0.7)

# ---------------------------------------------------------------- 13 IA
s = new_slide("Relação com IA generativa")
bullets(s, ["Carregar *embeddings*/*checkpoints*: `read()` = memória anônima; `mmap` = paginação sob demanda (e *thrashing* se não cabe)",
            "1ª carga fria ≫ próximas (cache de páginas)",
            "Força bruta O(N) → índices aproximados; contexto = custo **e** qualidade",
            "Filas limitadas = *backpressure*; caches de resposta só pagam com repetição"],
        0.6, 1.7, 12.1, 4.8, 22)

# ---------------------------------------------------------------- 14 limitações
s = new_slide("Limitações")
bullets(s, ["Contêiner sobre WSL2: \"frio\" = frio no cache da VM · 2 núcleos físicos, com ruído",
            "Geração e embeddings **simulados** → não generaliza para GPUs, modelos grandes ou ambiente distribuído",
            "Corrida com **injeção de falha** (janelas ampliadas) · swap e starvation não testados",
            "Causa total da saturação com mais de 2 workers não isolada (OpenBLAS explica ~10–15%)"],
        0.6, 1.7, 12.1, 4.8, 22)

# ---------------------------------------------------------------- 15 encerramento
s = new_slide("Obrigada", dark=True)
s.shapes.title.top, s.shapes.title.height = Inches(2.4), Inches(1.2)
s.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(48)
bullets(s, ["Repositório, relatório, dados, logs e instruções em `README.md`",
            "Reprodução: `python scripts/run_all.py`"], 0.6, 3.8, 12, 1.5, 24, "CADCFC", "Encerramento", bullet=False)

prs.save(OUT)
print("salvo:", OUT, "slides:", len(prs.slides))
